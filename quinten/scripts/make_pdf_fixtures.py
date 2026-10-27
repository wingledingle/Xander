#!/usr/bin/env python3
"""Builds fixed-layout test PDFs with LibreOffice: running headers, page-number footers,
a watermark line, headings, first-line-indent paragraphs and a two-column variant.
Text: chapters I-II of Lewis Carroll's 'Alice's Adventures in Wonderland' (public domain).
Requires python-docx and soffice. Output: core/analysis/src/test/resources/fixtures."""
import os, re, subprocess, sys, tempfile, urllib.request
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "core/analysis/src/test/resources/fixtures")
SRC = os.path.join(ROOT, "core/formats/build/external-fixtures/alice.txt")

def chapters():
    if not os.path.exists(SRC):
        os.makedirs(os.path.dirname(SRC), exist_ok=True)
        urllib.request.urlretrieve("https://www.gutenberg.org/cache/epub/11/pg11.txt", SRC)
    t = open(SRC, encoding="utf-8").read().replace("\r\n", "\n")
    start = t.index("CHAPTER I.\nDown the Rabbit-Hole")
    end = t.index("CHAPTER III.\nA Caucus-Race")
    body = t[start:end]
    parts = re.split(r"\n(?=CHAPTER [IVX]+\.\n)", body)
    out = []
    for p in parts:
        lines = p.strip().split("\n", 2)
        title = lines[0] + " " + lines[1]
        paras = [re.sub(r"\s+", " ", x).strip() for x in lines[2].split("\n\n")]
        out.append((title, [x for x in paras if x and not x.startswith("*")]))
    return out

def soft_hyphenate(s):
    def h(m):
        w = m.group(0)
        return w[:len(w)//2] + "­" + w[len(w)//2:] if len(w) >= 8 else w
    return re.sub(r"[A-Za-z]{8,}", h, s)

def field(par, instr):
    r = par.add_run()
    for tag, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instr
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        r._r.append(el)
        if tag == "separate":
            t = OxmlElement("w:t"); t.text = "1"; r._r.append(t)

def build(name, columns, indent):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(14.8), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    if columns:
        cols = OxmlElement("w:cols"); cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "500")
        sec._sectPr.append(cols)
    hp = sec.header.paragraphs[0]
    hp.text = "ALICE'S ADVENTURES IN WONDERLAND"
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.runs[0].font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(fp, "PAGE")
    wm = sec.footer.add_paragraph("Downloaded from free-example-books.test")
    wm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    wm.runs[0].font.size = Pt(6)
    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.font.name = "Liberation Serif"
    expected = []
    for ci, (title, paras) in enumerate(chapters()):
        h = doc.add_heading(title, level=1)
        if ci > 0:
            h.paragraph_format.page_break_before = True
        expected.append("#" + title)
        for p in paras:
            par = doc.add_paragraph(soft_hyphenate(p))
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = par.paragraph_format
            if indent:
                pf.first_line_indent = Cm(0.6); pf.space_after = Pt(0)
            else:
                pf.space_after = Pt(6)
            pf.line_spacing = 1.0
            expected.append(p)
    tmp = tempfile.mkdtemp()
    docx = os.path.join(tmp, name + ".docx")
    doc.save(docx)
    subprocess.check_call(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, docx], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    os.replace(os.path.join(tmp, name + ".pdf"), os.path.join(OUT, name + ".pdf"))
    with open(os.path.join(OUT, name + ".txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(expected) + "\n")

os.makedirs(OUT, exist_ok=True)
build("indent", columns=False, indent=True)
build("columns", columns=True, indent=False)
print("ok")
